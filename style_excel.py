# -*- coding: utf-8 -*-
"""
Script định dạng file Excel đẹp mắt, chuyên nghiệp:
- Đóng băng hàng tiêu đề (Freeze panes dòng 1)
- Header màu xanh Navy (Deep Blue #1E3A8A), chữ trắng đậm
- Tự động bật Wrap Text (xuống dòng) cho cột Yêu cầu và Biện pháp
- Căn chỉnh độ rộng cột và chiều cao hàng phù hợp
- Đóng khung viền mỏng tinh tế (#CBD5E1)
- Hiệu ứng Zebra striping (dòng xen kẽ màu nền xám sáng nhẹ #F8FAFC)
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def format_excel_file(file_path):
    if not os.path.exists(file_path):
        print(f"Không tìm thấy file: {file_path}")
        return False

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # 1. Đóng băng tiêu đề (dòng 1)
    ws.freeze_panes = "A2"

    # Định nghĩa màu sắc & font chữ
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue sang trọng
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Xám nhẹ xen kẽ
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    data_font = Font(name="Segoe UI", size=10, color="1E293B")

    thin_border = Side(border_style="thin", color="CBD5E1")
    border = Border(top=thin_border, left=thin_border, right=thin_border, bottom=thin_border)

    # 2. Định dạng hàng tiêu đề (Row 1)
    ws.row_dimensions[1].height = 34
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    # 3. Định dạng các hàng dữ liệu
    for row_idx in range(2, ws.max_row + 1):
        is_odd = (row_idx % 2 == 1)
        row_fill = zebra_fill if is_odd else white_fill
        row_cells = ws[row_idx]

        # Tính toán chiều cao hàng tự động dựa vào độ dài text
        yc_val = str(row_cells[2].value or "") if len(row_cells) > 2 else ""
        bp_val = str(row_cells[3].value or "") if len(row_cells) > 3 else ""

        lines_yc = max(1, len(yc_val) // 48 + 1)
        lines_bp = max(1, len(bp_val) // 48 + 1, bp_val.count("\n") + 1)
        max_lines = max(lines_yc, lines_bp)
        ws.row_dimensions[row_idx].height = max(45, max_lines * 19 + 10)

        for col_idx, cell in enumerate(row_cells):
            cell.font = data_font
            cell.fill = row_fill
            cell.border = border

            if col_idx == 0:
                # Cột STT
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx == 1:
                # Cột Tên / Nguồn ảnh
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                # Cột Nội dung yêu cầu & Biện pháp xử lý
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # 4. Căn chỉnh độ rộng các cột tối ưu
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 58

    wb.save(file_path)
    print(f"[OK] Đã định dạng thẩm mỹ và đóng băng tiêu đề cho file: {file_path}")
    return True

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    files = [
        os.path.join(base_dir, "Ket_Qua_Mau.xlsx"),
        os.path.join(base_dir, "Ket_Qua_Tong_Hop.xlsx"),
        os.path.join(base_dir, "Ket_Qua_Mau_Da_Sua_Xuong_Dong.xlsx"),
        os.path.join(base_dir, "Ket_Qua_Tong_Hop_Da_Sua_Xuong_Dong.xlsx")
    ]
    for f in files:
        if os.path.exists(f):
            format_excel_file(f)
