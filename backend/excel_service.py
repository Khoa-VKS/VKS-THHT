# -*- coding: utf-8 -*-
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ExcelService:
    @staticmethod
    def create_workbook(records):
        """
        records: danh sách dict gồm:
        [
            {"stt": 1, "name": "Ảnh 1", "yeucau": "...", "bienphap": "..."},
            ...
        ]
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lỗi Thường Gặp QLAHS"

        # 1. ĐÓNG BĂNG TIÊU ĐỀ HÀNG 1 KHI CUỘN CHUỘT
        ws.freeze_panes = "A2"

        # 2. Header
        headers = ["STT", "TÊN / NGUỒN ẢNH", "NỘI DUNG YÊU CẦU", "BIỆN PHÁP XỬ LÝ"]
        ws.append(headers)

        # 3. Ghi dữ liệu
        for idx, r in enumerate(records, 1):
            stt = r.get("stt", idx)
            name = r.get("name", f"Ảnh {idx}")
            yeucau = r.get("yeucau", "")
            bienphap = r.get("bienphap", "")
            ws.append([stt, name, yeucau, bienphap])

        # 4. Định dạng thẩm mỹ chuẩn
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        data_font = Font(name="Segoe UI", size=10, color="1E293B")

        thin = Side(border_style="thin", color="CBD5E1")
        border = Border(top=thin, left=thin, right=thin, bottom=thin)

        # Style hàng tiêu đề
        ws.row_dimensions[1].height = 34
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        # Style từng hàng dữ liệu
        for r_idx in range(2, len(records) + 2):
            is_odd = (r_idx % 2 == 1)
            row_fill = zebra_fill if is_odd else white_fill
            row_cells = ws[r_idx]

            # Tính chiều cao hàng tự động dựa theo độ dài chữ
            yc_len = len(str(row_cells[2].value or ""))
            bp_len = len(str(row_cells[3].value or ""))
            bp_newlines = str(row_cells[3].value or "").count("\n")
            max_lines = max(yc_len // 46 + 1, bp_len // 46 + 1, bp_newlines + 1)
            ws.row_dimensions[r_idx].height = max(50, max_lines * 19 + 10)

            for c_idx, cell in enumerate(row_cells):
                cell.font = data_font
                cell.fill = row_fill
                cell.border = border
                if c_idx == 0:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                elif c_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Độ rộng các cột
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 60
        ws.column_dimensions["D"].width = 60

        return wb

    @staticmethod
    def export_bytes(records):
        wb = ExcelService.create_workbook(records)
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def save_file(records, filepath):
        wb = ExcelService.create_workbook(records)
        wb.save(filepath)
        return filepath
