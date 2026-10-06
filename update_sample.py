# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

out_file = r"C:\Users\VKS\.gemini\antigravity\scratch\tach_noi_dung_excel\Ket_Qua_Mau.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Tổng Hợp Xử Lý"

# Đóng băng tiêu đề dòng 1
ws.freeze_panes = "A2"

headers = ["STT", "TÊN / NGUỒN ẢNH", "NỘI DUNG YÊU CẦU", "BIỆN PHÁP XỬ LÝ"]
ws.append(headers)

rows = [
    [
        1,
        "Vụ án 1 (Điều 251)",
        "Vụ án và bị can khởi tố theo khoản 2 Điều 251 là tội rất nghiêm trọng, thời hạn truy tố là 29 ngày nhưng sau khi chuyển sang giai đoạn truy tố chỉ có thời hạn 19 ngày.",
        "1. Không kiểm tra, điều chỉnh loại tội trước khi chuyển giai đoạn truy tố. Hiện vụ án vẫn loại tội NT.\n2. Chỉ liên hệ khi vụ án trên nền tảng đã quá hạn."
    ],
    [
        2,
        "Vụ án 2 (Nguyễn Xuân Thủy)",
        "Vụ án đã chuyển đến VKS Thành phố giải quyết theo thẩm quyền nhưng Viện kiểm sát thành phố nhận được vụ án nhưng không nhận được biện pháp ngăn chặn của bị can. Đề nghị chuyển biện pháp ngăn chặn của bị can Nguyễn Xuân Thủy cho VKS Thành phố để ra VKS Thành phố ban hành lệnh giam.",
        "1. Việc xác định tình trạng yêu cầu \"Rất khẩn\" chưa phù hợp với tính chất của hỗ trợ khi kiểm tra (bị can có thời hạn tạm giam đến ngày 19/8/2026).\n2. Vụ án sau khi Phòng 2 tiếp nhận đã tự \"Thêm mới\" 01 trường hợp tạm giam khác. Ghi nhận để nhắc nhở tại Hội nghị toàn quốc."
    ],
    [
        3,
        "Vụ án 3 (Trần Thanh An)",
        "Vụ án Trần Thanh An và đồng phạm đã nhập với vụ án khác nhưng biện pháp ngăn chặn vẫn ở lại tài khoản cũ, kính nhờ tổ công tác xóa giúp biện pháp ngăn chặn 04 bị can có hình ảnh kèm theo. Trân trọng!",
        "Vụ án tiến hành nhập ngày 06/8/2026. Sau khi hoàn thành nhập, không kiểm tra thông tin phần BPNC và tự thêm mới BPNCCC khác cho 04 bị can. Khi quá hạn tạm giam thì mới có yêu cầu kiểm tra xử lý dữ liệu. Nội dung này cần đưa vào vấn đề rút kinh nghiệm trong thời gian gần nhất."
    ]
]

for r in rows:
    ws.append(r)

header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
data_font = Font(name="Segoe UI", size=10, color="1E293B")

thin = Side(border_style="thin", color="CBD5E1")
border = Border(top=thin, left=thin, right=thin, bottom=thin)

ws.row_dimensions[1].height = 34
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

for r_idx in range(2, len(rows) + 2):
    is_odd = (r_idx % 2 == 1)
    row_fill = zebra_fill if is_odd else white_fill
    row_cells = ws[r_idx]
    
    # Auto height
    yc_len = len(str(row_cells[2].value or ""))
    bp_len = len(str(row_cells[3].value or ""))
    max_lines = max(yc_len // 48 + 1, bp_len // 48 + 1, str(row_cells[3].value or "").count("\n") + 1)
    ws.row_dimensions[r_idx].height = max(50, max_lines * 19 + 10)

    for c_idx, cell in enumerate(row_cells):
        cell.font = data_font
        cell.fill = row_fill
        cell.border = border
        if c_idx == 0:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif c_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 60

wb.save(out_file)
print("Updated Ket_Qua_Mau.xlsx with 3 cases!")
