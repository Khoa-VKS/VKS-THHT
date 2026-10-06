# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

out_files = [
    r"C:\Users\VKS\.gemini\antigravity\scratch\tach_noi_dung_excel\TONG_HOP_LOI_THUONG_GAP_KHI_THAO_TAC_NEN_TANG_QUAN_LY_AN_HINH_SU.xlsx",
    r"C:\Users\VKS\.gemini\antigravity\scratch\tach_noi_dung_excel\Ket_Qua_Mau.xlsx"
]

for out_file in out_files:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lỗi Thường Gặp QLAHS"

    # Đóng băng hàng tiêu đề 1
    ws.freeze_panes = "A2"

    headers = ["STT", "TÊN / NGUỒN ẢNH", "NỘI DUNG YÊU CẦU", "BIỆN PHÁP XỬ LÝ"]
    ws.append(headers)

    rows = [
        [
            1,
            "Ảnh 1 (Khu vực 17)",
            "Vụ án VA_E01. 59. 016_2608_0184 đã chuyển thành công đến Viện kiểm sát Khu vực 17, Thành phố Hồ Chí Minh giải quyết theo thẩm quyền. Trong vụ án có áp dụng biện pháp tạm giam đối với bị can Nguyễn Thành Kiên, tuy nhiên biện pháp ngăn chặn vẫn còn ở lại trên phần mềm Viện kiểm sát Khu vực 16, vậy kính nhờ bộ phận hỗ trợ QLAHS chuyển phần biện pháp ngăn chặn của bị can qua Viện kiểm sát Khu vực 17, Thành phố Hồ Chí Minh để kiểm sát theo thẩm quyền.",
            "1. Sau khi tách vụ án, người dùng không kiểm tra BPNC đã được tách cùng bị can hay chưa.\n2. Đã thực hiện chuyển vụ án khi chưa kiểm tra phần BPNC trong vụ án sau khi tách.\n3. Khu vực 17 đã hoàn thành luôn việc nhập vụ án."
        ],
        [
            2,
            "Ảnh 2 (Cục 2 hỗ trợ)",
            "Sau khi trình đề xuất, lãnh đạo không duyệt được đề xuất dẫn đến không thể chuyển giai đoạn. Kính mong Cục 2 xem xét và hỗ trợ.",
            "Văn bản phát sinh ngày 02/7/2026. Đến nay 02/10/2026 mới đề nghị hỗ trợ. Mã vụ án cung cấp không đúng."
        ],
        [
            3,
            "Ảnh 3 (Gia hạn thời hạn)",
            "Trong giai đoạn truy tố, không thao tác thêm mới \"Gia hạn thời hạn quyết định việc truy tố\" được. Dẫn đến báo quá hạn.",
            "Vụ án vừa gia hạn thời hạn truy tố ngày 08/8/2026. Chưa phát sinh việc trả ĐĐTBS và hoàn lại VKS. Không kiểm tra dữ liệu và thủ tục tố tụng trước khi đề nghị."
        ]
    ]

    for r in rows:
        ws.append(r)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10, color="1E293B")

    thin = Side(border_style="thin", color="CBD5E1")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    ws.row_dimensions[1].height = 34
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for r_idx in range(2, len(rows) + 2):
        is_odd = (r_idx % 2 == 1)
        row_fill = zebra_fill if is_odd else white_fill
        row_cells = ws[r_idx]
        
        yc_len = len(str(row_cells[2].value or ""))
        bp_len = len(str(row_cells[3].value or ""))
        max_lines = max(yc_len // 48 + 1, bp_len // 48 + 1, str(row_cells[3].value or "").count("\n") + 1)
        ws.row_dimensions[r_idx].height = max(50, max_lines * 19 + 10)

        for c_idx, cell in enumerate(row_cells):
            cell.font = data_font
            cell.fill = row_fill
            cell.border = border
            if c_idx == 0:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 60

    wb.save(out_file)
    print("Saved:", out_file)
