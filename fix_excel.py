# -*- coding: utf-8 -*-
"""
Script xử lý làm sạch văn bản, loại bỏ triệt để:
1. Mọi phần tiêu đề chụp dư bên trên ('Diễn biến xử lý...', 'NỘI DUNG YÊU CẦU', 'Biện pháp xử lý')
2. Mọi icon bình luận, đường kẻ ngang chụp dư ở đáy ảnh
3. Các dấu xuống dòng thừa giữa các câu (nối câu liền mạch)
"""

import os
import re
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def clean_yeu_cau(text):
    """
    Chỉ lấy phần nội dung nằm bên dưới 'NỘI DUNG YÊU CẦU' và bên trên 'Biện pháp xử lý'.
    Loại bỏ hoàn toàn 'Diễn biến xử lý...', 'NỘI DUNG YÊU CẦU'.
    """
    if not text or not str(text).strip():
        return ""
    text = str(text).strip()

    # 1. Nếu vô tình dính cả phần Biện pháp xử lý, cắt bỏ phần từ Biện pháp xử lý trở đi
    bp_match = re.search(r'\n?[^\n]*(?:biện pháp xử lý|bien phap xu ly)[^\n]*\n?', text, re.IGNORECASE)
    if bp_match:
        text = text[:bp_match.start()].strip()

    # 2. Tìm điểm kết thúc của dòng tiêu đề NỘI DUNG YÊU CẦU
    # Cắt bỏ mọi thứ đứng trước và bao gồm cả cụm từ NỘI DUNG YÊU CẦU
    yc_match = re.search(r'(?:.*)(?:nội dung yêu cầu|noi dung yeu cau)\s*[:\-\.]?\s*', text, re.IGNORECASE | re.DOTALL)
    if yc_match:
        text = text[yc_match.end():].strip()
    else:
        # Nếu không có từ NỘI DUNG YÊU CẦU nhưng có Diễn biến xử lý...
        text = re.sub(r'^[^\n]*(?:diễn biến xử lý|dien bien xu ly)[^\n]*\s*', '', text, flags=re.IGNORECASE).strip()

    # 3. Nối câu liền mạch
    return unwrap_text(text, is_numbered=False)

def clean_bien_phap(text):
    """
    Chỉ lấy phần nội dung nằm bên dưới tiêu đề màu xanh 'Biện pháp xử lý'.
    Loại bỏ chữ 'Biện pháp xử lý', loại bỏ icon bình luận [💬] và đường kẻ rác ở đáy.
    """
    if not text or not str(text).strip():
        return ""
    text = str(text).strip()

    # 1. Cắt bỏ chữ 'Biện pháp xử lý' nếu có ở đầu
    bp_match = re.search(r'^(?:[^\n]*)(?:biện pháp xử lý|bien phap xu ly)\s*[:\-\.]?\s*', text, re.IGNORECASE)
    if bp_match:
        text = text[bp_match.end():].strip()

    # 2. Bỏ các dòng rác ở đáy ảnh (icon bình luận, đường kẻ, ký tự đơn lẻ)
    lines = text.split('\n')
    while lines:
        last = lines[-1].strip()
        # Nếu dòng cuối rất ngắn hoặc chỉ chứa ký tự đặc biệt, icon chat
        if len(last) <= 3 or last in ['---', '___', '===', 'Q', 'O', '[]', '[ ]', '💬']:
            lines.pop()
        else:
            break
    text = '\n'.join(lines).strip()

    # 3. Nối liền mạch hoặc ngắt dòng sạch giữa các mục 1., 2.
    return unwrap_text(text, is_numbered=True)

def unwrap_text(text, is_numbered=False):
    """
    Nối các dòng thuộc cùng 1 câu/đoạn bằng dấu cách.
    Chỉ giữ lại xuống dòng khi bắt đầu mục mới (1., 2., -, +...).
    """
    if not text or not str(text).strip():
        return ""
    text = str(text)

    # Tách các mục số bị dính liền (vd: "hạn.2. Chỉ liên hệ" -> "hạn.\n2. Chỉ liên hệ")
    text = re.sub(r'([\.\?\!\;])\s*(\d+[\.\)])\s*', r'\1\n\2 ', text)

    raw_lines = [l.strip() for l in text.splitlines()]
    raw_lines = [l for l in raw_lines if l]
    if not raw_lines:
        return ""

    bullet_regex = re.compile(r'^(\d+[\.\/\)]|\-|\+|\*|•|[a-zA-Z]\))\s*')

    items = []
    current_item = ""

    for line in raw_lines:
        # Lọc bỏ nếu còn sót tiêu đề
        lower = line.lower()
        if "diễn biến xử lý" in lower or "dien bien xu ly" in lower:
            continue
        if lower in ["nội dung yêu cầu", "noi dung yeu cau", "biện pháp xử lý", "bien phap xu ly"]:
            continue

        is_bullet = bool(bullet_regex.match(line))

        if is_numbered and is_bullet:
            if current_item:
                items.append(current_item.strip())
            current_item = line
        else:
            if not current_item:
                current_item = line
            else:
                if current_item.endswith("-"):
                    current_item = current_item[:-1] + line
                else:
                    current_item += " " + line

    if current_item:
        items.append(current_item.strip())

    if not is_numbered and len(items) > 1 and not any(bullet_regex.match(it) for it in items):
        joined = " ".join(items)
    else:
        joined = "\n".join(items)

    joined = re.sub(r'[ \t]+', ' ', joined)
    joined = re.sub(r' *\n *', '\n', joined).strip()
    return joined

def clean_excel_file(input_file, output_file=None):
    if not os.path.exists(input_file):
        print(f"Lỗi: Không tìm thấy file {input_file}")
        return False

    if not output_file:
        base, ext = os.path.splitext(input_file)
        output_file = f"{base}_Da_Loc_Sach{ext}"

    wb = openpyxl.load_workbook(input_file)
    ws = wb.active

    header_row = 1
    yc_col = 3
    bp_col = 4

    for col_idx in range(1, ws.max_column + 1):
        val = str(ws.cell(row=header_row, column=col_idx).value or "").strip().upper()
        if "YÊU CẦU" in val or "YEU CAU" in val:
            yc_col = col_idx
        elif "BIỆN PHÁP" in val or "BIEN PHAP" in val or "XỬ LÝ" in val or "XU LY" in val:
            bp_col = col_idx

    modified_count = 0
    for r in range(header_row + 1, ws.max_row + 1):
        if yc_col:
            cell_yc = ws.cell(row=r, column=yc_col)
            if cell_yc.value:
                cleaned = clean_yeu_cau(cell_yc.value)
                if cleaned != cell_yc.value:
                    cell_yc.value = cleaned
                    modified_count += 1

        if bp_col:
            cell_bp = ws.cell(row=r, column=bp_col)
            if cell_bp.value:
                cleaned = clean_bien_phap(cell_bp.value)
                if cleaned != cell_bp.value:
                    cell_bp.value = cleaned
                    modified_count += 1

    wb.save(output_file)
    print(f"[OK] Đã lọc sạch {modified_count} ô dữ liệu (loại bỏ chữ chụp dư và icon đáy)!")
    return True

if __name__ == "__main__":
    if len(sys.argv) > 1:
        target = sys.argv[1]
        out = sys.argv[2] if len(sys.argv) > 2 else None
        clean_excel_file(target, out)
